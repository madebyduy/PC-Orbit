from __future__ import annotations

from pathlib import Path
from typing import Iterable
from zipfile import ZIP_DEFLATED, ZipFile
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Thu Ha\Desktop\PC Orbit")
REFERENCE = Path(
    r"C:\Users\Thu Ha\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.1"
    r"\skills\artifact-template-system-design\assets\reference.docx"
)
OUTPUT = ROOT / "docs" / "research" / "PC_Orbit_Deep_Research_Functions_2026-09-04.docx"

NAVY = "0C3154"
NAVY_2 = "163E62"
SLATE = "5A7189"
TEXT = "22364D"
BLUE = "2D6FA3"
PALE = "E4EFF7"
PALE_2 = "F3F7FA"
PALE_3 = "D4E6F3"
WHITE = "FFFFFF"
AMBER = "F4B544"
AMBER_PALE = "FFF3D6"
GREEN = "2F7D69"
GREEN_PALE = "E3F1ED"
RED = "B5473B"
RED_PALE = "FBE9E7"


SOURCES = {
    "S01": (
        "Microsoft — Recovery options in Windows",
        "https://support.microsoft.com/en-us/windows/experience/backup-recovery/recovery-options-in-windows",
    ),
    "S02": (
        "Microsoft Learn — Quick Machine Recovery",
        "https://learn.microsoft.com/en-us/windows/configuration/quick-machine-recovery/",
    ),
    "S03": (
        "Microsoft — Point-in-time restore for Windows",
        "https://support.microsoft.com/en-us/windows/experience/backup-recovery/point-time-restore-for-windows",
    ),
    "S04": (
        "Microsoft — Recovery Drive",
        "https://support.microsoft.com/en-us/windows/experience/backup-recovery/recovery-drive",
    ),
    "S05": (
        "Microsoft Learn — Windows RE technical reference",
        "https://learn.microsoft.com/en-us/windows-hardware/manufacture/desktop/windows-recovery-environment--windows-re--technical-reference?view=windows-11",
    ),
    "S06": (
        "Microsoft Learn — WinPE introduction",
        "https://learn.microsoft.com/en-us/windows-hardware/manufacture/desktop/winpe-intro?view=windows-11",
    ),
    "S07": (
        "Microsoft Learn — SetupDiag",
        "https://learn.microsoft.com/en-us/windows/deployment/upgrade/setupdiag",
    ),
    "S08": (
        "Microsoft Learn — BitLocker FAQ",
        "https://learn.microsoft.com/en-us/windows/security/operating-system-security/data-protection/bitlocker/faq",
    ),
    "S09": (
        "Microsoft Learn — BitLocker recovery known issues",
        "https://learn.microsoft.com/en-us/troubleshoot/windows-client/windows-security/bitlocker-recovery-known-issues",
    ),
    "S10": (
        "Microsoft Learn — Windows UEFI firmware update platform",
        "https://learn.microsoft.com/en-us/windows-hardware/drivers/bringup/windows-uefi-firmware-update-platform",
    ),
    "S11": (
        "UEFI Specification 2.11 — Firmware Update and Reporting",
        "https://uefi.org/specs/UEFI/2.11/23_Firmware_Update_and_Reporting.html",
    ),
    "S12": (
        "Dell — Command Configure: Manage BIOS settings",
        "https://www.dell.com/support/manuals/en-us/command-configure/dcc_ug_5.x/manage-bios-settings?guid=guid-5d7ed65b-4b6a-4236-9c91-b16f0e5b20e1",
    ),
    "S13": (
        "Lenovo — BIOS WMI Interface Guide",
        "https://docs.lenovocdrt.com/ref/bios/wmi/wmi_guide/",
    ),
    "S14": (
        "HP — Client Management Script Library",
        "https://h10032.www1.hp.com/ctg/Manual/c06696094.pdf",
    ),
    "S15": (
        "Microsoft Learn — PnPUtil command syntax",
        "https://learn.microsoft.com/en-us/windows-hardware/drivers/devtest/pnputil-command-syntax",
    ),
    "S16": (
        "Microsoft Learn — Driver selection process",
        "https://learn.microsoft.com/en-us/windows-hardware/drivers/install/overview-of-the-driver-selection-process",
    ),
    "S17": (
        "Microsoft — Update drivers through Device Manager",
        "https://support.microsoft.com/en-us/windows/update-drivers-through-device-manager-in-windows-ec62f46c-ff14-c91d-eead-d7126dc1f7b6",
    ),
    "S18": (
        "Microsoft Learn — Add/remove drivers to an offline image",
        "https://learn.microsoft.com/en-us/windows-hardware/manufacture/desktop/add-and-remove-drivers-to-an-offline-windows-image?view=windows-11",
    ),
    "S19": (
        "Microsoft Learn — WHEA hardware error events",
        "https://learn.microsoft.com/en-us/windows-hardware/drivers/whea/whea-hardware-error-events",
    ),
    "S20": (
        "Microsoft Learn — Analyze a kernel-mode dump with WinDbg",
        "https://learn.microsoft.com/en-us/windows-hardware/drivers/debugger/analyzing-a-kernel-mode-dump-file-with-windbg",
    ),
    "S21": (
        "Microsoft Learn — Windows diagnostic data privacy",
        "https://learn.microsoft.com/en-us/windows/privacy/configure-windows-diagnostic-data-in-your-organization",
    ),
    "S22": (
        "Microsoft Learn — powercfg command-line options",
        "https://learn.microsoft.com/en-us/windows-hardware/design/device-experiences/powercfg-command-line-options",
    ),
    "S23": (
        "Microsoft — Windows Backup",
        "https://support.microsoft.com/en-us/windows/experience/backup-recovery/back-up-and-restore-with-windows-backup",
    ),
    "S24": (
        "Microsoft Learn — winget export",
        "https://learn.microsoft.com/en-au/windows/package-manager/winget/export",
    ),
    "S25": (
        "Microsoft Learn — WSL basic commands",
        "https://learn.microsoft.com/en-us/windows/wsl/basic-commands",
    ),
    "S26": (
        "AMD — How to update motherboard BIOS",
        "https://www.amd.com/en/resources/support-articles/faqs/cpu-99.html",
    ),
    "S27": (
        "ASUS — CPU/Memory QVL lookup",
        "https://www.asus.com/support/FAQ/1043883",
    ),
    "S28": (
        "MSI — Boot / no display troubleshooting",
        "https://www.msi.com/support/technical_details/MB_Boot_No_Display",
    ),
    "S29": (
        "Microsoft Learn — UCSI",
        "https://learn.microsoft.com/en-us/windows-hardware/drivers/usbcon/ucsi",
    ),
    "S30": (
        "Microsoft Learn — USBView",
        "https://learn.microsoft.com/en-us/windows-hardware/drivers/debugger/usbview",
    ),
    "S31": (
        "Microsoft Learn — QueryDisplayConfig",
        "https://learn.microsoft.com/en-us/windows/win32/api/winuser/nf-winuser-querydisplayconfig",
    ),
    "S32": (
        "Microsoft Learn — Secure boot",
        "https://learn.microsoft.com/en-us/windows-hardware/design/device-experiences/oem-secure-boot",
    ),
    "S33": (
        "Microsoft — Device security in Windows Security",
        "https://support.microsoft.com/en-us/windows/security/windows-security/device-security-in-the-windows-security-app",
    ),
    "S34": (
        "Microsoft Learn — Windows security baselines",
        "https://learn.microsoft.com/en-us/windows/security/operating-system-security/device-management/windows-security-configuration-framework/windows-security-baselines",
    ),
    "S35": (
        "Microsoft — Quick Assist",
        "https://support.microsoft.com/en-us/windows/apps/solve-pc-problems-remotely-using-quick-assist",
    ),
    "S36": (
        "Microsoft — Reset your PC",
        "https://support.microsoft.com/en-us/windows/experience/backup-recovery/reset-your-pc",
    ),
    "S37": (
        "NIST SP 800-88 Rev.2 — Guidelines for Media Sanitization",
        "https://csrc.nist.gov/pubs/sp/800/88/r2/final",
    ),
    "S38": (
        "Microsoft — Free up drive space in Windows",
        "https://support.microsoft.com/en-us/windows/experience/storage-filemanagement/free-up-drive-space-in-windows",
    ),
    "S39": (
        "Microsoft — Storage settings in Windows",
        "https://support.microsoft.com/en-us/windows/experience/storage-filemanagement/storage-settings-in-windows",
    ),
    "S40": (
        "Microsoft Sysinternals — Autoruns for Windows",
        "https://learn.microsoft.com/en-us/sysinternals/downloads/autoruns",
    ),
    "S41": (
        "Microsoft Learn — Windows Performance Recorder",
        "https://learn.microsoft.com/en-us/windows-hardware/test/wpt/windows-performance-recorder",
    ),
}


FEATURES = [
    {
        "rank": 1,
        "phase": "P0",
        "score": "4,8",
        "name": "Recovery Readiness & Emergency Kit",
        "job": "Biết chắc máy có thể tự cứu trước khi sự cố xảy ra.",
        "functions": [
            "Kiểm tra WinRE, recovery media, quyền truy cập khóa BitLocker, backup gần nhất, dung lượng và khả năng boot UEFI/USB.",
            "Đóng gói driver mạng/lưu trữ, manifest phần cứng, edition/activation, checksum media và đường recovery OEM theo model.",
            "Recovery drill không phá hủy; QR/case code để mở hướng dẫn trên điện thoại khi PC không hiển thị.",
        ],
        "guard": "Không lưu recovery key/secret vào kit; không coi Recovery Drive là backup dữ liệu.",
        "sources": ["S04", "S05", "S08"],
    },
    {
        "rank": 2,
        "phase": "P0",
        "score": "4,7",
        "name": "Install / Reinstall Concierge",
        "job": "Cài hoặc sửa Windows mà không xóa nhầm ổ, mất dữ liệu, mạng, driver hay kích hoạt.",
        "functions": [
            "Phân loại đúng repair, reinstall giữ app/file/settings, reset hay clean install; ước lượng mất mát và downtime.",
            "Migration manifest: Windows Backup, winget, WSL, app settings, license cần xử lý thủ công.",
            "Nhận diện ổ bằng model/serial/dung lượng/sơ đồ phân vùng; chuẩn bị driver bootstrap; first-boot agent tiếp tục plan và hậu kiểm.",
        ],
        "guard": "Không tự chọn/xóa partition khi identity mơ hồ; không nhúng password/recovery key vào unattend.",
        "sources": ["S01", "S07", "S23", "S24", "S25"],
    },
    {
        "rank": 3,
        "phase": "P0",
        "score": "4,6",
        "name": "Root-cause Timeline",
        "job": "Trả lời điều gì thay đổi ngay trước khi máy bắt đầu lỗi.",
        "functions": [
            "Hợp nhất update, app, PnP/driver, firmware, WER/minidump, WHEA, Setup logs và thay đổi do PC Orbit.",
            "Xếp hạng giả thuyết theo thời gian, subsystem match, mức lặp lại và khả năng đảo ngược.",
            "Mỗi giả thuyết có bằng chứng ủng hộ/chống, confidence, phép thử tiếp theo và điều kiện dừng.",
        ],
        "guard": "Chỉ nói tương quan/khả năng khi chưa đủ chứng cứ nhân quả; dump là opt-in riêng.",
        "sources": ["S07", "S19", "S20", "S21"],
    },
    {
        "rank": 4,
        "phase": "P0",
        "score": "4,5",
        "name": "Boot & Recovery Orchestrator",
        "job": "Dẫn người dùng qua lỗi không boot theo đường ít mất mát nhất.",
        "functions": [
            "Ladder: repair/reinstall-in-place → gỡ update/driver → restore → WinRE/Startup Repair → reset → recovery media/clean install.",
            "Nhận diện Quick Machine Recovery trên Windows/build được hỗ trợ và khả năng mạng trong WinRE.",
            "Safe Mode/offline driver rollback, DISM→SFC theo ngữ cảnh; lưu kế hoạch bền vững qua reboot.",
        ],
        "guard": "Hiển thị rõ dữ liệu/app/settings sẽ mất; không lặp auto-repair vô hạn.",
        "sources": ["S01", "S02", "S03", "S05", "S18"],
    },
    {
        "rank": 5,
        "phase": "P0",
        "score": "4,5",
        "name": "Driver Steward",
        "job": "Cài đúng driver cho đúng hardware ID và quay lại được khi driver mới gây lỗi.",
        "functions": [
            "Inventory provider/signer/rank/date/version/package/problem code; Windows Update trước, OEM chính thức khi cần.",
            "Driver Vault và bootstrap pack; cài từng subsystem; health window và rollback.",
            "Giải thích xung đột memory integrity/vulnerable-driver policy trước khi đề xuất tắt bảo vệ.",
        ],
        "guard": "Không update-all, không nguồn driver tổng hợp, không ép unsigned hoặc gỡ boot-critical mù quáng.",
        "sources": ["S15", "S16", "S17", "S18", "S33"],
    },
    {
        "rank": 6,
        "phase": "P0",
        "score": "4,4",
        "name": "BIOS/UEFI Safe Navigator + Firmware Guardian",
        "job": "Thay đổi/cập nhật firmware đúng model mà không khóa BitLocker hoặc làm máy không boot.",
        "functions": [
            "Map capability chuẩn sang tên/menu đúng model, revision và BIOS version; ảnh từng bước + phone mode.",
            "Bốn mức Auto / Guided / Read-only / Unsupported; preflight AC/pin, BitLocker, password, chữ ký, reboot count, recovery.",
            "Chỉ dùng UEFI capsule/ESRT hoặc công cụ OEM chính thức; hậu kiểm version/outcome và resume BitLocker.",
            "Theo dõi chuyển đổi chứng chỉ Secure Boot 2023 do chứng chỉ 2011 bắt đầu hết hạn từ 06/2026.",
        ],
        "guard": "Không raw NVRAM write; model mismatch hoặc knowledge pack hết hạn phải fail closed.",
        "sources": ["S08", "S09", "S10", "S11", "S12", "S13", "S14", "S32"],
    },
    {
        "rank": 7,
        "phase": "P1",
        "score": "4,4",
        "name": "Privacy-safe Support Bundle & Handoff",
        "job": "Gửi đủ dữ liệu cho người hỗ trợ nhưng không gửi nhầm thông tin riêng tư.",
        "functions": [
            "Bundle theo sự cố: summary, timeline, model/OS/build/driver/firmware và test đã chạy.",
            "Preview/redact username, path, IP/MAC, SSID, serial, token, key; dump cần consent riêng.",
            "Case code + Quick Assist với cảnh báo lừa đảo, quyền điều khiển rõ và log phiên.",
        ],
        "guard": "Mặc định local-only; không upload trước khi người dùng duyệt nội dung.",
        "sources": ["S21", "S35"],
    },
    {
        "rank": 8,
        "phase": "P1",
        "score": "4,2",
        "name": "Hardware Health & Degradation",
        "job": "Phân biệt linh kiện xuống cấp với lỗi phần mềm bằng xu hướng, không bằng một health score giả.",
        "functions": [
            "NVMe/storage errors, wear, temperature, unsafe shutdown; WHEA corrected/nonfatal/fatal theo thời gian.",
            "Battery design/full capacity, sleep drain, energy/sleepstudy, wake blockers và throttling khi có telemetry.",
            "Baseline theo chính máy; tách rõ không có lỗi, không hỗ trợ telemetry và chưa đủ thời gian quan sát.",
        ],
        "guard": "Không tuyên bố linh kiện tốt chỉ vì API trả trống; không gán phần trăm sức khỏe thiếu cơ sở.",
        "sources": ["S19", "S22"],
    },
    {
        "rank": 9,
        "phase": "P1",
        "score": "4,1",
        "name": "New Build & No-POST Copilot",
        "job": "Hướng dẫn tự lắp/nâng cấp khi bấm nguồn không lên hoặc không có hình.",
        "functions": [
            "Trước mua: socket/chipset, CPU support + BIOS floor, RAM QVL, lane sharing, nguồn/đầu cắm và kích thước khi có nguồn chính thức.",
            "Assembly map đúng manual; POST triage theo fan/LED/beep/Q-LED/Q-code; minimal POST, một DIMM, monitor path, clear CMOS.",
            "BIOS Flashback/OEM recovery đúng model; hướng dẫn offline trên điện thoại.",
        ],
        "guard": "Không dùng sơ đồ bo khác revision; mọi bước chạm phần cứng phải có power-off/ESD warning phù hợp.",
        "sources": ["S26", "S27", "S28"],
    },
    {
        "rank": 10,
        "phase": "P1",
        "score": "4,1",
        "name": "Security Compatibility Planner",
        "job": "Bật Secure Boot/TPM/BitLocker/memory integrity cho mục tiêu thật mà không phá tương thích.",
        "functions": [
            "Outcome graph cho Windows 11/game/bảo mật; phân biệt unsupported, disabled, driver-blocked, firmware-needed và Unknown.",
            "Kiểm tra driver, firmware, UEFI/GPT và recovery trước; cập nhật phụ thuộc rồi mới bật protection.",
            "Readiness cho Secure Boot certificate transition; không áp enterprise baseline máy móc cho consumer.",
        ],
        "guard": "Tắt protection chỉ là ngoại lệ tạm thời, có thời hạn và phương án khôi phục.",
        "sources": ["S32", "S33", "S34"],
    },
    {
        "rank": 11,
        "phase": "P2",
        "score": "3,8",
        "name": "Migration Blueprint",
        "job": "Chuyển sang PC mới và biết chính xác thứ gì đã hoặc chưa được khôi phục.",
        "functions": [
            "Hợp nhất Windows Backup, winget, WSL, app-setting vault, peripheral profiles và activation checklist.",
            "Reconciliation desired vs observed; dry-run; app/license nào cần người dùng đăng nhập hoặc cấu hình lại.",
        ],
        "guard": "Secret không vào blueprint; artifact nhạy cảm phải mã hóa và có retention rõ.",
        "sources": ["S23", "S24", "S25"],
    },
    {
        "rank": 12,
        "phase": "P2",
        "score": "3,8",
        "name": "Upgrade Compatibility Planner",
        "job": "Biết nâng CPU/RAM/SSD/GPU có chạy thật và phải chuẩn bị gì trước khi tháo máy.",
        "functions": [
            "Compatibility graph: CPU support/BIOS floor, QVL, slots/lanes, PCIe, power/connector/physical constraints.",
            "Kế hoạch thứ tự: firmware với CPU cũ → recovery → lắp → minimal POST → verify → tối ưu sau.",
        ],
        "guard": "Mọi kết luận có URL OEM/ngày truy xuất; QVL là tested list, không phải danh sách duy nhất có thể chạy.",
        "sources": ["S26", "S27"],
    },
    {
        "rank": 13,
        "phase": "P2",
        "score": "3,6",
        "name": "Connection Path & Port Advisor",
        "job": "Tìm đúng nút thắt cổng/cáp/dock khi USB chậm, sạc yếu hoặc mất màn hình.",
        "functions": [
            "Vẽ controller → hub/dock → device; descriptor, negotiated speed/power và slow-charge status.",
            "Vẽ display path/connector/mode; hướng dẫn A/B test một mắt xích và timed auto-revert cho display.",
        ],
        "guard": "Không kết luận do cáp nếu không có phép thử hoặc telemetry hỗ trợ.",
        "sources": ["S29", "S30", "S31"],
    },
    {
        "rank": 14,
        "phase": "P2",
        "score": "3,6",
        "name": "Device Retirement Wizard",
        "job": "Bán/tặng/thải bỏ PC mà dữ liệu không còn và tài khoản được xử lý đúng.",
        "functions": [
            "Backup/restore check, sign-out/unlink, encryption/recovery, reset/clean theo mục tiêu.",
            "Với dữ liệu nhạy cảm: phương pháp sanitize + validation phù hợp media và NIST SP 800-88 Rev.2.",
        ],
        "guard": "Không gọi consumer reset là chứng nhận sanitization; chỉ xuất biên bản đúng với bằng chứng đã có.",
        "sources": ["S36", "S37"],
    },
]


CLAIMS = [
    ("C01", "Recovery phải đi từ ít gián đoạn đến phá hủy hơn; reinstall qua Windows Update có thể giữ dữ liệu/app/settings.", ["S01"]),
    ("C02", "Quick Machine Recovery dùng WinRE có mạng để tìm remediation và chỉ là best-effort trên build phù hợp.", ["S02"]),
    ("C03", "Point-in-time restore có thể hoàn nguyên file/app/settings gần đây và cần khóa BitLocker khi được yêu cầu.", ["S03"]),
    ("C04", "Recovery Drive không chứa file cá nhân và cần được tạo lại định kỳ.", ["S04"]),
    ("C05", "BitLocker có thể vào recovery sau thay đổi firmware/TPM/boot; update nhiều reboot cần suspension đúng.", ["S08", "S09"]),
    ("C06", "Windows có nền tảng UEFI capsule/ESRT; firmware rollback không đồng nhất giữa platform.", ["S10", "S11"]),
    ("C07", "BIOS APIs và điều kiện password/model khác nhau giữa Dell, Lenovo và HP.", ["S12", "S13", "S14"]),
    ("C08", "PnPUtil và driver ranking cho phép inventory/export có căn cứ; không nên dùng version lớn nhất làm chuẩn duy nhất.", ["S15", "S16", "S17"]),
    ("C09", "Offline servicing driver có thể sửa boot nhưng cũng có thể làm image không boot nếu gỡ driver trọng yếu.", ["S18"]),
    ("C10", "SetupDiag, WHEA và dump cung cấp bằng chứng khác nhau; dump có rủi ro riêng tư.", ["S07", "S19", "S20", "S21"]),
    ("C11", "Windows Backup, winget và WSL chỉ là các mảnh của migration blueprint.", ["S23", "S24", "S25"]),
    ("C12", "CPU mới có thể cần BIOS floor; QVL và no-POST steps là dữ liệu theo motherboard/model.", ["S26", "S27", "S28"]),
    ("C13", "UCSI/USBView/DisplayConfig cho phép xây dựng topology cổng và display path.", ["S29", "S30", "S31"]),
    ("C14", "Chứng chỉ Secure Boot 2011 bắt đầu hết hạn từ 06/2026; readiness phải feature-detect và theo nguồn hiện hành.", ["S32"]),
    ("C15", "Tắt memory integrity làm giảm bảo mật; security baseline cần cân nhắc tác động vận hành.", ["S33", "S34"]),
    ("C16", "Quick Assist cần consent và cảnh báo lừa đảo; support handoff phải minh bạch quyền.", ["S35"]),
    ("C17", "Reset clean-data không phải chuẩn sanitization cho ngành/chính phủ; NIST Rev.2 là chuẩn tham chiếu hiện hành.", ["S36", "S37"]),
    ("C18", "Windows có các nhóm dọn dung lượng chính thức như temporary/system files, large or unused files, cloud-synced files và unused apps; xóa Windows.old làm mất khả năng quay lại phiên bản trước.", ["S38", "S39"]),
    ("C19", "Autoruns quan sát nhiều điểm tự khởi động như services, scheduled tasks, shell extensions và logon entries; dữ liệu này phù hợp cho inventory trước khi vô hiệu hóa có kiểm soát.", ["S40"]),
    ("C20", "Windows Performance Recorder thu thập ETW trace để phân tích mức tiêu thụ tài nguyên và hiệu năng bằng Windows Performance Analyzer.", ["S41"]),
]


EXPANSION_GROUPS = [
    (
        "Ứng dụng, dọn rác và hiệu năng",
        [
            ("Startup & Background Controller", "Chỉ ra app/service/task nào làm chậm boot, đăng nhập hoặc chạy nền; vô hiệu hóa từng mục có thể hoàn tác.", "P0 · Guided"),
            ("Performance Root-cause Profiler", "Thu ETW theo kịch bản CPU, disk, memory, boot, sleep hoặc ứng dụng rồi diễn giải nguyên nhân có bằng chứng.", "P0 · Observe"),
            ("Why is my PC slow? Investigator", "Tách chậm do CPU, RAM, pagefile, storage, nhiệt, update, antivirus, startup hay mạng thay vì đề xuất tối ưu chung.", "P0 · Observe"),
            ("Safe Software Installer", "Kiểm chữ ký, publisher, kiến trúc, nguồn tải, yêu cầu quyền và tạo checkpoint trước khi cài.", "P1 · Guided"),
            ("Runtime & Dependency Doctor", "Tìm thiếu hoặc xung đột .NET, VC++ runtime, DirectX, WebView, Java, Python hay biến môi trường theo ứng dụng.", "P1 · Guided"),
            ("Trial Installation Sandbox", "Cho thử ứng dụng trong môi trường cách ly hoặc snapshot, quan sát thay đổi rồi giữ hay loại bỏ.", "P1 · Guided"),
            ("Application Conflict Detector", "Tương quan crash/hang với overlay, hook, antivirus, driver, plugin và ứng dụng vừa cài.", "P1 · Observe"),
            ("Default Apps Repair", "Khôi phục file association, protocol handler và hành vi mở file mà không ghi registry mù quáng.", "P1 · Guided"),
            ("Broken Uninstaller Rescue", "Gỡ app hỏng dựa trên manifest, package, service, task và startup entry; preview phần còn lại.", "P1 · Guided"),
            ("Update Regression Guard", "Ghi baseline trước update, theo dõi lỗi sau update và đề xuất rollback đúng package/driver khi đủ bằng chứng.", "P0 · Observe"),
        ],
    ),
    (
        "Dữ liệu, ổ đĩa và di chuyển",
        [
            ("Partition Safety Assistant", "Mô phỏng thao tác phân vùng, nhận diện recovery/EFI/OEM và chặn xóa khi mục tiêu còn mơ hồ.", "P0 · Guided"),
            ("Disk Clone & SSD Migration", "Kiểm dung lượng, BitLocker, sector size, boot mode và hậu kiểm boot sau khi clone sang SSD.", "P1 · Guided"),
            ("Data Rescue Mode", "Ưu tiên image/read-only copy, đánh giá sức khỏe ổ và tránh ghi thêm lên thiết bị nghi hỏng.", "P1 · Guided"),
            ("Folder Relocation Manager", "Di chuyển thư mục người dùng/game/library có preview, kiểm quyền, link và đường quay lại.", "P1 · Guided"),
            ("Cloud Sync Conflict Resolver", "Phân biệt local-only, online-only, conflict copy và deletion propagation trước khi dọn dữ liệu.", "P1 · Observe"),
            ("File Organization Assistant", "Nhóm file theo loại, tuổi, dự án và mức sử dụng; mọi đổi tên/di chuyển đều có dry-run và undo.", "P2 · Guided"),
            ("Sensitive File Finder", "Tìm file có nguy cơ chứa khóa, giấy tờ, bản sao dữ liệu hoặc secret mà không tự upload nội dung.", "P1 · Observe"),
            ("Backup Integrity Monitor", "Không chỉ báo lần backup cuối; thử đọc, kiểm manifest/hash mẫu và cảnh báo đích backup không còn truy cập.", "P0 · Observe"),
        ],
    ),
    (
        "Tài khoản, bảo mật và hỗ trợ từ xa",
        [
            ("Login Readiness Check", "Kiểm đường đăng nhập, PIN, mật khẩu, khóa khôi phục, tài khoản dự phòng và truy cập mạng trước thay đổi lớn.", "P0 · Observe"),
            ("User Profile Repair", "Phân biệt lỗi profile, quyền, sync và shell; tạo đường chuyển profile an toàn thay vì sửa trực tiếp dữ liệu mù quáng.", "P1 · Guided"),
            ("Admin Safety Manager", "Giảm chạy bằng admin thường xuyên, giải thích prompt quyền và tạo thao tác nâng quyền theo phiên có audit.", "P1 · Guided"),
            ("Account Migration", "Đối soát tài khoản cục bộ/Microsoft, profile, khóa mã hóa và quyền sở hữu khi đổi người dùng hoặc đổi máy.", "P1 · Guided"),
            ("Family / Senior Mode", "Giao diện ít lựa chọn, chữ rõ, diễn giải rủi ro, nút gọi trợ giúp và khóa thao tác phá hủy.", "P1 · UX"),
            ("Remote-support Shield", "Xác minh người hỗ trợ, hiển thị quyền đang cấp, cảnh báo lừa đảo, ghi phiên và nút ngắt khẩn cấp.", "P0 · Guided"),
            ("Suspicious Persistence Scanner", "Inventory autorun, service, task, extension và thay đổi bất thường; không tự kết luận mọi mục lạ là malware.", "P1 · Observe"),
            ("Post-malware Recovery", "Kiểm persistence, tài khoản, proxy/DNS, Defender state, file hệ thống và kế hoạch đổi credential sau sự cố.", "P1 · Guided"),
            ("Ransomware Readiness", "Kiểm backup tách biệt, khả năng restore, quyền thư mục và playbook cô lập máy trước khi có sự cố.", "P1 · Observe"),
            ("Download Reputation & Signature", "Kiểm publisher/chữ ký/hash/nguồn, gắn nhãn unknown và giải thích rủi ro trước khi chạy file.", "P1 · Observe"),
        ],
    ),
    (
        "Firmware và phần cứng chuyên sâu",
        [
            ("BIOS Baseline & Diff", "Chụp cấu hình BIOS đọc được, so sánh trước/sau reset/update và diễn giải setting nào thực sự thay đổi.", "P0 · Observe"),
            ("CMOS Reset Recovery Plan", "Chuẩn bị ảnh cấu hình, BitLocker, boot order và profile cần phục hồi trước khi clear CMOS.", "P0 · Guided"),
            ("Firmware Dependency Chain", "Biểu diễn thứ tự BIOS, EC, dock, SSD, GPU firmware và driver; chặn chuỗi cập nhật không tương thích.", "P1 · Observe"),
            ("RAM Stability Assistant", "Theo dõi XMP/EXPO, lỗi bộ nhớ/WHEA, nhiệt và phép thử từng DIMM/slot với một biến mỗi lần.", "P1 · Guided"),
            ("PSU & Power Event Investigator", "Tương quan reboot/mất nguồn với tải, WHEA, sự kiện điện và thay đổi phần cứng; tránh khẳng định PSU hỏng quá sớm.", "P1 · Observe"),
            ("Noise Source Finder", "Phân biệt quạt, coil whine, HDD, pump và rung cộng hưởng qua bài thử tải có giới hạn.", "P2 · Guided"),
            ("Dust & Maintenance Planner", "Lập lịch vệ sinh theo nhiệt/độ ồn/môi trường và hướng dẫn tắt nguồn, ESD, tháo lắp an toàn.", "P2 · Guided"),
            ("Hardware Change Timeline", "Ghi lần thay RAM/SSD/GPU/cáp/dock và firmware để nối sự cố với thay đổi vật lý.", "P1 · Observe"),
            ("Manual Vault", "Lưu manual, sơ đồ header, recovery procedure và checksum theo đúng model/revision để dùng offline.", "P1 · Observe"),
        ],
    ),
    (
        "Màn hình, âm thanh và thiết bị ngoại vi",
        [
            ("Display Safe Setup", "Thử độ phân giải, tần số, HDR, GPU path và tự quay lại nếu mất tín hiệu.", "P0 · Guided"),
            ("Multi-monitor Workspace Restore", "Lưu topology, scale, orientation, primary display và vị trí cửa sổ theo dock/location.", "P1 · Guided"),
            ("Audio Routing Doctor", "Vẽ input/output, format, exclusive mode, enhancement, Bluetooth profile và app routing.", "P1 · Observe"),
            ("Webcam & Meeting Check", "Kiểm camera, microphone, permission, ánh sáng, echo và network trước cuộc họp.", "P1 · Observe"),
            ("Printer Queue Surgeon", "Tách lỗi queue, spooler, driver, port, discovery và mạng; dọn job kẹt có xác nhận.", "P1 · Guided"),
            ("Bluetooth Reliability Manager", "Theo dõi adapter, profile, interference, power saving và pairing history; thử từng thay đổi có rollback.", "P1 · Guided"),
            ("Gaming & Creator Stability Lab", "Tạo workload tái hiện crash/stutter, theo dõi frame-time, nhiệt, driver, overlay và thiết bị capture.", "P1 · Observe"),
            ("Workload Profiles", "Cấu hình theo mục tiêu game, học, họp, render hoặc tiết kiệm pin với before/after và nút hoàn nguyên.", "P2 · Guided"),
            ("Energy Cost Tracker", "Ước lượng điện năng theo workload/thời gian và phát hiện sleep drain hoặc máy không thật sự idle.", "P2 · Observe"),
        ],
    ),
    (
        "Chẩn đoán khó, xác minh và vòng đời",
        [
            ("Freeze / Hang Recorder", "Bắt trace vòng đệm và ảnh chụp trạng thái khi UI treo; lưu mốc trước/sau thay vì chờ người dùng mô tả.", "P0 · Observe"),
            ("Intermittent Fault Diary", "Ghi lỗi chập chờn theo thời gian, nguồn điện, nhiệt, thiết bị gắn ngoài, mạng và thao tác gần nhất.", "P0 · Observe"),
            ("Known-issue Matcher", "Đối chiếu model/build/driver/firmware với known issue còn hiệu lực và nêu rõ độ chắc chắn.", "P1 · Observe"),
            ("A/B Experiment Manager", "Thiết kế phép thử một biến, thời lượng, metric, điều kiện dừng và tự phục hồi cấu hình.", "P0 · Guided"),
            ("Second-opinion Mode", "So sánh nhiều giả thuyết/nguồn và chỉ ra bằng chứng còn thiếu trước khi chấp nhận một sửa chữa rủi ro.", "P1 · Observe"),
            ("Repair Verification Suite", "Chạy lại triệu chứng, health window và kiểm tra độc lập sau sửa; không báo thành công chỉ vì lệnh chạy xong.", "P0 · Observe"),
            ("What can my PC do? Advisor", "Ghép phần cứng/phần mềm với workload thực, nêu nút thắt và nâng cấp có lợi thay vì điểm số chung.", "P2 · Observe"),
            ("Warranty, License & Recall Monitor", "Lưu bằng chứng mua, serial cục bộ, license, hạn bảo hành và cảnh báo recall/firmware chính thức.", "P2 · Observe"),
            ("PC Handover Package", "Tạo gói bàn giao máy gồm inventory, trạng thái, việc đã làm, dữ liệu đã xóa và phần còn cần người nhận xử lý.", "P2 · Guided"),
            ("Offline Phone Companion", "Đưa checklist, QR, ảnh và recovery plan sang điện thoại khi PC không boot, không có mạng hoặc đang ở BIOS.", "P1 · UX"),
        ],
    ),
]


EXPANSION_PRIORITIES = [
    ("1", "Smart Cleanup", "Tần suất cao, giá trị thấy ngay; chỉ Auto cho dữ liệu có quy tắc an toàn và có quarantine/undo."),
    ("2", "Freeze / Hang Recorder", "Thu được bằng chứng ở lỗi khó tái hiện; bổ trợ trực tiếp Root-cause Timeline."),
    ("3", "Startup & Background Controller", "Giải quyết boot/login chậm nhưng vẫn giữ nguyên tắc một thay đổi mỗi lần."),
    ("4", "Performance Root-cause Profiler", "Biến cảm giác “máy chậm” thành trace và subsystem cụ thể."),
    ("5", "App Repair + Broken Uninstaller", "Khoảng trống thường gặp sau cài/gỡ app; có thể triển khai Guided trước."),
    ("6", "Update Regression Guard", "Giảm sự cố sau update bằng baseline, timeline và rollback đúng đối tượng."),
    ("7", "BIOS Baseline & Diff", "Giảm rủi ro khi reset/update BIOS và hỗ trợ chẩn đoán model-specific."),
    ("8", "Trial Installation Sandbox", "Hạn chế tác động app lạ; cần lựa chọn công nghệ cách ly phù hợp edition."),
    ("9", "Partition Safety + SSD Migration", "Giá trị cao nhưng cần gate nghiêm, test boot và nhận diện ổ tuyệt đối rõ."),
    ("10", "Backup Integrity Monitor", "Đóng vòng recovery-readiness bằng khả năng đọc/khôi phục thực, không chỉ timestamp."),
]


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_borders(table, color="FFFFFF", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_cell_width(cell, width_inches: float) -> None:
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, text: str, url: str, color=BLUE, underline=True):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")
    r_fonts.set(qn("w:eastAsia"), "Arial")
    r_pr.append(r_fonts)
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])


def clear_document_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def tune_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = rgb(TEXT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    pf = normal.paragraph_format
    pf.space_after = Pt(5)
    pf.line_spacing = 1.15

    for style_name, size, color, before, after in (
        ("Heading 1", 20, NAVY, 10, 7),
        ("Heading 2", 14, NAVY, 10, 4),
        ("Heading 3", 11.5, SLATE, 8, 3),
    ):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(36)
    title.font.bold = True
    title.font.color.rgb = rgb(NAVY)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    title.paragraph_format.space_after = Pt(4)

    if "PCO Small" not in [s.name for s in doc.styles]:
        small = doc.styles.add_style("PCO Small", 1)
    else:
        small = doc.styles["PCO Small"]
    small.font.name = "Arial"
    small.font.size = Pt(8.3)
    small.font.color.rgb = rgb(SLATE)
    small._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    small.paragraph_format.space_after = Pt(2)


def style_run(run, *, size=None, color=None, bold=None, italic=None) -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_body(doc: Document, text: str = "", *, bold=False, italic=False, color=None, after=5, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_together = keep
    r = p.add_run(text)
    style_run(r, bold=bold, italic=italic, color=color)
    return p


def add_bullet(doc: Document, text: str, level: int = 0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.22 + level * 0.18)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    marker = p.add_run("•  ")
    style_run(marker, bold=True, color=BLUE)
    p.add_run(text)
    return p


def add_numbered(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    marker = p.add_run("›  ")
    style_run(marker, bold=True, color=BLUE, size=12)
    p.add_run(text)
    return p


def add_citations(doc: Document, ids: Iterable[str]):
    ids = list(ids)
    p = doc.add_paragraph(style="PCO Small")
    p.paragraph_format.keep_together = True
    r = p.add_run("Căn cứ: ")
    style_run(r, bold=True, color=SLATE)
    for idx, source_id in enumerate(ids):
        title, url = SOURCES[source_id]
        add_hyperlink(p, f"[{source_id}]", url, color=BLUE, underline=False)
        if idx != len(ids) - 1:
            p.add_run("  ")
    return p


def add_callout(doc: Document, title: str, body: str, *, kind="info"):
    palette = {
        "info": (PALE, NAVY),
        "decision": (GREEN_PALE, GREEN),
        "warning": (AMBER_PALE, NAVY),
        "risk": (RED_PALE, RED),
    }
    fill, accent = palette[kind]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    prevent_row_split(table.rows[0])
    set_cell_width(table.cell(0, 0), 0.13)
    set_cell_width(table.cell(0, 1), 6.83)
    set_cell_shading(table.cell(0, 0), accent)
    set_cell_shading(table.cell(0, 1), fill)
    set_table_borders(table, color=fill, size="0")
    left = table.cell(0, 0)
    right = table.cell(0, 1)
    set_cell_margins(left, 0, 0, 0, 0)
    set_cell_margins(right, 110, 150, 110, 150)
    p = right.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    style_run(r, bold=True, color=accent, size=10.5)
    p2 = right.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.add_run(body)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def format_table(table, header=True, first_col=False, widths=None, font_size=9.0):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color=WHITE, size="5")
    for r_idx, row in enumerate(table.rows):
        prevent_row_split(row)
        if r_idx == 0 and header:
            set_repeat_table_header(row)
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                set_cell_width(cell, widths[c_idx])
            if header and r_idx == 0:
                set_cell_shading(cell, NAVY)
            elif first_col and c_idx == 0:
                set_cell_shading(cell, PALE_3)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, PALE_2)
            else:
                set_cell_shading(cell, WHITE)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for run in p.runs:
                    style_run(run, size=font_size, color=WHITE if (header and r_idx == 0) else TEXT, bold=(header and r_idx == 0))


def fill_cell(cell, text: str, *, bold=False, color=None, size=9.0, clear=True):
    if clear:
        cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    style_run(r, bold=bold, color=color, size=size)
    return p


def add_section_title(doc: Document, number: str, title: str, *, page_break=False):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.keep_with_next = True
    p.add_run(f"{number}.  {title}")
    return p


def add_phase_badge(paragraph, phase: str):
    colors = {"P0": RED, "P1": AMBER, "P2": BLUE}
    r = paragraph.add_run(f"  {phase}")
    style_run(r, bold=True, color=colors[phase], size=9.5)


def configure_footer(doc: Document) -> None:
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    doc.settings.odd_and_even_pages_header_footer = True
    for footer in (section.footer, section.even_page_footer):
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.clear()
        r = p.add_run("PC ORBIT  |  DEEP RESEARCH  •  04.09.2026  |  ")
        style_run(r, color=SLATE, size=8)
        add_page_field(p)
    first = section.first_page_footer
    first.paragraphs[0].clear()
    section.header.paragraphs[0].clear()
    section.first_page_header.paragraphs[0].clear()


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Inches(2.05)
    spacer.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("PC Orbit")
    style_run(r, size=38, color=SLATE, bold=False)

    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(6)
    p.add_run("Trợ lý tự phục vụ PC an toàn")
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Inches(1.08)
    r2 = p2.add_run("Nghiên cứu chức năng chuyên sâu — cài máy, BIOS/UEFI, chẩn đoán, cứu hộ và toàn bộ vòng đời thiết bị")
    style_run(r2, size=15, color=TEXT, bold=True)

    meta = doc.add_table(rows=2, cols=3)
    meta.autofit = False
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = ["TRẠNG THÁI", "CHỦ SỞ HỮU", "CẬP NHẬT LẦN CUỐI"]
    values = ["Proposed", "PC Orbit", "04/09/2026"]
    for idx in range(3):
        set_cell_width(meta.cell(0, idx), 2.32)
        set_cell_width(meta.cell(1, idx), 2.32)
        set_cell_margins(meta.cell(0, idx), 20, 60, 20, 60)
        set_cell_margins(meta.cell(1, idx), 20, 60, 60, 60)
        p0 = fill_cell(meta.cell(0, idx), labels[idx], bold=True, color=SLATE, size=8.5)
        p1 = fill_cell(meta.cell(1, idx), values[idx], color=TEXT, size=10.2)
        p0.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
    set_table_borders(meta, color=WHITE, size="0")

    details = doc.add_table(rows=4, cols=2)
    details.autofit = False
    details.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels2 = ["Tác giả", "Độc giả", "Tài liệu liên quan", "Phạm vi"]
    values2 = [
        "OpenAI — Deep Research cho PC Orbit",
        "Product, Engineering, UX, QA và đối tác phần cứng/OEM",
        "README.md; PC_CONTROL_PROJECT_SPEC_v3.md; PC_CONTROL_ARCHITECTURE.md",
        "Windows PC cá nhân: trước mua/lắp, cài đặt, vận hành, BIOS/UEFI, sửa lỗi, phục hồi, chuyển máy và thanh lý",
    ]
    for i in range(4):
        fill_cell(details.cell(i, 0), labels2[i], bold=True, color=NAVY, size=9.2)
        fill_cell(details.cell(i, 1), values2[i], color=TEXT, size=9.2)
    format_table(details, header=False, first_col=True, widths=[1.35, 5.61], font_size=9.2)

    doc.add_page_break()


def add_executive_summary(doc: Document) -> None:
    add_section_title(doc, "1", "Tóm tắt điều hành")
    add_body(
        doc,
        "PC Orbit không nên trở thành một bảng có thật nhiều nút chỉnh hay một chatbot đưa lệnh ngẫu nhiên. "
        "Khoảng trống lớn nhất là một hệ thống điều phối tự phục vụ an toàn: nhận diện đúng máy, đọc trạng thái thật, "
        "chọn phương án ít phá hủy nhất, chuẩn bị đường lui, tiếp tục qua khởi động lại và xác minh bằng nguồn độc lập.",
        after=7,
    )
    add_callout(
        doc,
        "Quyết định đề xuất",
        "Định vị PC Orbit là “PC self-service safety system”. Ưu tiên Recovery Readiness, Install/Reinstall Concierge, "
        "Root-cause Timeline, Boot Recovery, Driver Steward và BIOS/UEFI Safe Navigator trước khi mở rộng sang tối ưu hóa.",
        kind="decision",
    )
    add_body(doc, "Sáu luồng P0", bold=True, color=SLATE, after=3)
    for item in FEATURES[:6]:
        p = add_numbered(doc, f"{item['name']}: {item['job']}")
        p.paragraph_format.keep_together = True
    add_body(
        doc,
        "Ba yếu tố tạo khác biệt thực sự: model-specific thay vì hướng dẫn chung; evidence-first thay vì đoán lỗi; "
        "recovery-first thay vì chỉ tối ưu lúc máy đang chạy.",
        bold=True,
        color=NAVY,
        after=6,
    )
    add_body(
        doc,
        "Lớp mở rộng nên bắt đầu bằng Smart Cleanup có phân vùng an toàn và hoàn tác, Startup/Background Controller, "
        "Performance Profiler, App Repair, Update Regression Guard, Freeze Recorder và BIOS Baseline & Diff.",
        italic=True,
        color=SLATE,
        after=5,
    )
    add_citations(doc, ["S01", "S08", "S10", "S15"])


def add_scope_goals(doc: Document) -> None:
    add_section_title(doc, "2", "Phạm vi, giả định và phương pháp")
    add_body(
        doc,
        "Nghiên cứu bao phủ Windows 10/11 trên laptop/máy bộ OEM và desktop tự lắp, từ chọn linh kiện đến thanh lý. "
        "Đầu vào gồm tài liệu Microsoft, UEFI, NIST và OEM chính thức, chốt ngày 04/09/2026; đồng thời đối chiếu với "
        "kiến trúc và đặc tả hiện có của PC Orbit.",
    )
    add_body(doc, "Cách đọc kết quả", bold=True, color=SLATE, after=2)
    add_bullet(doc, "Mức P0/P1/P2 là khuyến nghị lộ trình; điểm ưu tiên là ước lượng chuyên gia, chưa phải số liệu thị trường.")
    add_bullet(doc, "Auto chỉ hợp lệ khi có adapter chính thức, test phần cứng thật và đường recovery; nếu không phải Guided/Read-only/Unsupported.")
    add_bullet(doc, "Thông tin phụ thuộc thời điểm như Windows recovery và Secure Boot phải feature-detect lúc chạy.")

    add_section_title(doc, "3", "Mục tiêu và không-mục-tiêu")
    table = doc.add_table(rows=5, cols=2)
    goals = [
        "Người phổ thông hoàn tất luồng rủi ro cao mà không cần terminal.",
        "Mọi thay đổi đều có preflight, preview, verify và đường phục hồi.",
        "Hỗ trợ BIOS/driver/phần cứng theo model và mức tin cậy có kiểm chứng.",
        "Handoff cho kỹ thuật viên đủ bằng chứng nhưng bảo vệ riêng tư.",
    ]
    non_goals = [
        "Universal BIOS flasher hoặc ghi raw UEFI/NVRAM.",
        "Kho “repair scripts” không có state/dependency/verify.",
        "AI tự quyết định/xác nhận thay người dùng cho thao tác phá hủy.",
        "Cam kết mọi máy hoặc mọi linh kiện đều có telemetry/API.",
    ]
    fill_cell(table.cell(0, 0), "Mục tiêu", bold=True, color=WHITE, size=9.5)
    fill_cell(table.cell(0, 1), "Không-mục-tiêu", bold=True, color=WHITE, size=9.5)
    for i in range(4):
        fill_cell(table.cell(i + 1, 0), goals[i], size=9.3)
        fill_cell(table.cell(i + 1, 1), non_goals[i], size=9.3)
    format_table(table, header=True, widths=[3.48, 3.48], font_size=9.3)


def add_problem_and_lifecycle(doc: Document) -> None:
    add_section_title(doc, "4", "Vấn đề người dùng và vòng đời PC")
    add_body(
        doc,
        "Người dùng hiếm khi thiếu một lệnh; họ thiếu sự chắc chắn rằng lệnh đó đúng với đúng máy, có làm mất dữ liệu, "
        "có kích hoạt BitLocker recovery, có cần driver mạng/lưu trữ, và có quay lại được không. Các công cụ Windows/OEM "
        "hiện có mạnh nhưng rời rạc. PC Orbit nên nối chúng thành hành trình có ngữ cảnh và trạng thái bền vững.",
    )
    rows = [
        ("1", "Chọn / nâng cấp", "Tương thích, BIOS floor, QVL, nguồn/cổng/kích thước", "Mua nhầm hoặc máy không POST"),
        ("2", "Lắp / POST", "Sơ đồ đúng model, debug LED/code, minimal POST", "Cắm sai, không hình, flash sai"),
        ("3", "Cài / chuyển máy", "Backup, media, disk identity, driver bootstrap, activation", "Xóa nhầm ổ, mất mạng/dữ liệu"),
        ("4", "Vận hành", "Baseline, health, power, security compatibility", "Tối ưu mù, tắt bảo vệ"),
        ("5", "Sự cố", "Timeline, giả thuyết, một phép thử/lần", "Random fixes, không biết nguyên nhân"),
        ("6", "Không boot", "WinRE, rollback, recovery ladder", "Reset/clean quá sớm"),
        ("7", "Handoff / thanh lý", "Redaction, remote consent, sanitization evidence", "Rò dữ liệu hoặc xóa chưa đủ"),
    ]
    table = doc.add_table(rows=1 + len(rows), cols=4)
    for j, h in enumerate(["Giai đoạn", "Việc người dùng", "Năng lực cần có", "Rủi ro chính"]):
        fill_cell(table.cell(0, j), h, bold=True, color=WHITE, size=8.4)
    for i, row in enumerate(rows, start=1):
        for j, value in enumerate(row):
            fill_cell(table.cell(i, j), value, size=8.1)
    format_table(table, header=True, first_col=True, widths=[0.65, 1.35, 3.05, 1.91], font_size=8.1)


def add_architecture(doc: Document) -> None:
    add_section_title(doc, "5", "Kiến trúc năng lực đề xuất", page_break=True)
    add_body(
        doc,
        "Lõi hiện có của PC Orbit là nền phù hợp. Phần mở rộng nên giữ một invariant: không side effect rủi ro khi "
        "chưa khóa được device identity, evidence, policy gate, recovery path và verification contract.",
    )

    banner = doc.add_table(rows=1, cols=1)
    cell = banner.cell(0, 0)
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell, 100, 140, 100, 140)
    p = fill_cell(cell, "PC ORBIT — SAFE SELF-SERVICE CONTROL PLANE", bold=True, color=WHITE, size=10)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r = p2.add_run("Ý định → bằng chứng → chính sách → kế hoạch → thực thi/guided → phục hồi → xác minh")
    style_run(r, color=PALE_3, size=8.5)
    set_table_borders(banner, color=NAVY, size="0")

    flow = doc.add_table(rows=2, cols=4)
    labels = [
        ("1", "Ý định & UX", "Windows UI • Phone/QR"),
        ("2", "Device identity", "Model • revision • IDs"),
        ("3", "Evidence graph", "State • source • freshness"),
        ("4", "Planner & gates", "Dependencies • risk • preview"),
        ("5", "Adapters", "Windows • OEM • Guided"),
        ("6", "Recovery", "Checkpoint • WinRE • rollback"),
        ("7", "Verification", "Independent read • outcome"),
        ("8", "Timeline & handoff", "Audit • privacy • support"),
    ]
    for idx, (n, title, sub) in enumerate(labels):
        row, col = divmod(idx, 4)
        c = flow.cell(row, col)
        set_cell_shading(c, PALE if row == 0 else PALE_2)
        set_cell_margins(c, 110, 100, 110, 100)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(n)
        style_run(rr, bold=True, color=BLUE, size=13)
        p2 = c.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(1)
        rr2 = p2.add_run(title)
        style_run(rr2, bold=True, color=NAVY, size=8.8)
        p3 = c.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_after = Pt(0)
        rr3 = p3.add_run(sub)
        style_run(rr3, color=SLATE, size=7.4)
    format_table(flow, header=False, widths=[1.74, 1.74, 1.74, 1.74], font_size=8)
    doc.add_paragraph("Hình 1. Luồng điều khiển an toàn theo trạng thái; các ô là ranh giới trách nhiệm, không phải màn hình sản phẩm.", style="PCO Small")

    add_body(doc, "Các khối lõi", bold=True, color=SLATE, after=3)
    components = [
        ("Device identity & evidence", "Model/revision/hardware IDs; state, source, freshness, confidence", "Unknown; yêu cầu xác nhận; không side effect"),
        ("Outcome compiler", "Dependency graph, conflict, preflight, plan preview", "Không biên dịch khi thiếu điều kiện an toàn"),
        ("Safety policy engine", "BitLocker, nguồn, backup, quyền, mật khẩu, support tier", "Fail closed cho firmware write/xóa ổ"),
        ("Execution adapters", "Windows API/CLI allowlist, OEM API, capsule, guided", "Auto → Guided → Read-only → Unsupported"),
        ("Recovery coordinator", "Checkpoint, durable resume, timeout auto-revert, WinRE/USB", "Dừng an toàn; không retry vô hạn"),
        ("Verification & timeline", "Đọc lại độc lập, outcome, history, hypotheses, handoff", "Gắn nhãn chưa xác minh; không success giả"),
        ("Signed knowledge packs", "Manual/model map, QVL, nguồn chính thức, expiry", "Pack hết hạn: observe/guided only"),
    ]
    table = doc.add_table(rows=1 + len(components), cols=3)
    for j, h in enumerate(["Khối", "Trách nhiệm", "Fallback / failure behavior"]):
        fill_cell(table.cell(0, j), h, bold=True, color=WHITE, size=8.5)
    for i, row in enumerate(components, start=1):
        for j, value in enumerate(row):
            fill_cell(table.cell(i, j), value, size=8.2)
    format_table(table, header=True, first_col=True, widths=[1.55, 3.27, 2.14], font_size=8.2)


def add_safety_contracts(doc: Document) -> None:
    add_section_title(doc, "6", "Hợp đồng an toàn và vòng đời một yêu cầu")
    steps = [
        "Người dùng chọn outcome và phạm vi thiết bị; hệ thống đóng băng device fingerprint cho plan.",
        "Đọc evidence từ nguồn độc lập, kèm source/freshness/confidence; Unknown không bị ép thành false.",
        "Compiler dựng dependency, conflict, preflight, mức hỗ trợ và kế hoạch ít phá hủy nhất.",
        "Policy engine kiểm recovery path, BitLocker, nguồn, quyền, backup, model match và secret handling.",
        "Ghi durable checkpoint và approved-plan hash trước side effect.",
        "Thực thi qua adapter allowlist hoặc guided step; mỗi lần chỉ thay đổi một biến khi chẩn đoán.",
        "Qua reboot/WinRE, resume đúng plan; timeout hoặc mất tín hiệu dùng rollback/auto-revert đã định trước.",
        "Reader độc lập xác minh outcome; timeline ghi before/after, nguồn, lỗi, rollback và dữ liệu hỗ trợ.",
    ]
    for s in steps:
        add_numbered(doc, s)
    add_callout(
        doc,
        "Invariants không được phá",
        "Không firmware write/xóa ổ nếu device identity hoặc recovery path không chắc chắn. Không coi exit code là verify. "
        "Không upload log/dump trước khi preview và consent. Không để AI tạo lệnh thực thi ngoài capability allowlist.",
        kind="warning",
    )

    add_body(doc, "Hợp đồng dữ liệu tối thiểu", bold=True, color=SLATE, after=3)
    data_rows = [
        ("DeviceIdentity", "model, board_revision, firmware_version, hardware_ids", "Có", "Khóa đúng máy/adapter/knowledge pack"),
        ("Evidence", "value, source, observed_at, freshness, confidence", "Có", "Phân biệt fact, inference và Unknown"),
        ("RiskGate", "kind, status, rationale, blocking", "Có", "Không side effect nếu blocking gate fail"),
        ("RecoveryAsset", "type, location, verified_at, contains_secret", "Có khi rủi ro", "Đường cứu hộ và chính sách lưu trữ"),
        ("PlanStep", "adapter, preconditions, effect, verify, rollback", "Có", "Đơn vị giao dịch/replay"),
        ("Hypothesis", "evidence_for, against, confidence, next_test", "Chẩn đoán", "Giải thích nguyên nhân có kiểm soát"),
        ("SupportBundleManifest", "included, redacted, consent, retention", "Handoff", "Privacy boundary có audit"),
    ]
    table = doc.add_table(rows=1 + len(data_rows), cols=4)
    for j, h in enumerate(["Record", "Field chính", "Khi nào", "Guarantee"]):
        fill_cell(table.cell(0, j), h, bold=True, color=WHITE, size=8.2)
    for i, row in enumerate(data_rows, start=1):
        for j, value in enumerate(row):
            fill_cell(table.cell(i, j), value, size=7.9)
    format_table(table, header=True, first_col=True, widths=[1.35, 2.55, 1.05, 2.01], font_size=7.9)


def add_feature_catalog(doc: Document) -> None:
    add_section_title(doc, "7", "Danh mục chức năng chuyên sâu")
    add_body(
        doc,
        "Mỗi chức năng dưới đây được mô tả như một hệ thống hoàn chỉnh: việc người dùng cần giải quyết, năng lực lõi, "
        "rào chắn và nguồn kỹ thuật. P0 nên được xây trước; P1/P2 mở rộng khi nền an toàn đã đạt launch gate.",
    )
    for idx, feature in enumerate(FEATURES):
        p = doc.add_paragraph(style="Heading 2")
        p.paragraph_format.keep_with_next = True
        p.add_run(f"7.{idx + 1}  {feature['name']}")
        add_phase_badge(p, feature["phase"])

        mini = doc.add_table(rows=1, cols=3)
        vals = [
            ("ƯU TIÊN", f"#{feature['rank']}"),
            ("ĐIỂM ƯỚC LƯỢNG", f"{feature['score']} / 5"),
            ("JOB TO BE DONE", feature["job"]),
        ]
        widths = [0.82, 1.23, 4.91]
        for c_idx, (label, value) in enumerate(vals):
            cell = mini.cell(0, c_idx)
            set_cell_shading(cell, PALE if c_idx < 2 else PALE_2)
            set_cell_margins(cell, 75, 90, 75, 90)
            p1 = cell.paragraphs[0]
            p1.paragraph_format.space_after = Pt(1)
            r1 = p1.add_run(label)
            style_run(r1, bold=True, color=SLATE, size=7.1)
            p2 = cell.add_paragraph()
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(value)
            style_run(r2, bold=(c_idx < 2), color=NAVY, size=8.4)
        format_table(mini, header=False, widths=widths, font_size=8.3)

        for item in feature["functions"]:
            add_bullet(doc, item)
        add_callout(doc, "Rào chắn", feature["guard"], kind="risk" if feature["phase"] == "P0" else "warning")
        add_citations(doc, feature["sources"])


def add_expansion_catalog(doc: Document) -> None:
    add_section_title(doc, "8", "Hệ sinh thái chức năng mở rộng")
    add_body(
        doc,
        "Danh mục này mở rộng PC Orbit từ các luồng cứu hộ cốt lõi sang bảo trì hằng ngày, ứng dụng, dữ liệu, tài khoản, "
        "thiết bị ngoại vi và các lỗi chập chờn. Đây là backlog có cấu trúc: ưu tiên vẫn phụ thuộc nghiên cứu người dùng, "
        "khả năng quan sát thật của Windows/OEM và mức rủi ro của từng adapter.",
    )

    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.keep_with_next = True
    p.add_run("8.1  Smart Cleanup — dọn rác có bằng chứng và hoàn tác")
    add_phase_badge(p, "P0")
    add_body(
        doc,
        "Smart Cleanup không phải nút “xóa càng nhiều càng tốt”. Nó phải lập bản đồ dung lượng, giải thích vì sao một mục "
        "có thể xóa, ước lượng dung lượng thu hồi, cho xem trước và giữ đường hoàn tác.",
    )
    cleanup_rows = [
        ("Tự động an toàn", "Temporary/cache đã hết hạn, thumbnail cache, Recycle Bin quá hạn, log cũ theo retention đã biết", "Auto chỉ khi không bị ứng dụng khóa; ghi before/after"),
        ("Cần người dùng duyệt", "Downloads, file lớn/cũ, ứng dụng ít dùng, cloud offline files, crash dumps, cache game/browser", "Hiển thị owner, tuổi, vị trí, tác động và dung lượng"),
        ("Không xóa trực tiếp", "WinSxS, DriverStore, recovery partition, restore point, registry và app data không rõ", "Dùng API/công cụ được hỗ trợ hoặc hạ xuống hướng dẫn"),
    ]
    table = doc.add_table(rows=1 + len(cleanup_rows), cols=3)
    for j, h in enumerate(["Vùng", "Ví dụ", "Chính sách"]):
        fill_cell(table.cell(0, j), h, bold=True, color=WHITE, size=8.4)
    for i, row in enumerate(cleanup_rows, start=1):
        for j, value in enumerate(row):
            fill_cell(table.cell(i, j), value, size=8.0)
    format_table(table, header=True, first_col=True, widths=[1.28, 3.18, 2.50], font_size=8.0)
    for item in [
        "Storage map theo người dùng, ứng dụng và loại dữ liệu; tìm file trùng bằng hash và nhận diện download/installer chưa hoàn tất.",
        "Dọn phần còn lại của ứng dụng bằng manifest, service, scheduled task và startup entry; không suy đoán theo tên thư mục đơn thuần.",
        "Quarantine 7–30 ngày, exact reclaim estimate, post-clean verification và lịch sử phục hồi từng mục.",
        "Bảo vệ Windows.old khi người dùng còn cần rollback; không có registry cleaner, RAM cleaner hoặc dọn WinSxS/DriverStore mù quáng.",
    ]:
        add_bullet(doc, item)
    add_citations(doc, ["S38", "S39"])

    for group_idx, (group_name, rows) in enumerate(EXPANSION_GROUPS, start=2):
        p = doc.add_paragraph(style="Heading 2")
        p.paragraph_format.keep_with_next = True
        p.add_run(f"8.{group_idx}  {group_name}")
        table = doc.add_table(rows=1 + len(rows), cols=3)
        for j, h in enumerate(["Module", "Giá trị người dùng", "Mức đề xuất"]):
            fill_cell(table.cell(0, j), h, bold=True, color=WHITE, size=8.0)
        for i, row in enumerate(rows, start=1):
            for j, value in enumerate(row):
                fill_cell(table.cell(i, j), value, size=7.35)
        format_table(table, header=True, first_col=True, widths=[1.72, 4.15, 1.09], font_size=7.35)
        if group_idx == 2:
            add_citations(doc, ["S40", "S41"])

    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.keep_with_next = True
    p.add_run("8.8  Danh sách nên đưa vào discovery trước")
    table = doc.add_table(rows=1 + len(EXPANSION_PRIORITIES), cols=3)
    for j, h in enumerate(["#", "Cụm chức năng", "Lý do"]):
        fill_cell(table.cell(0, j), h, bold=True, color=WHITE, size=8.2)
    for i, row in enumerate(EXPANSION_PRIORITIES, start=1):
        for j, value in enumerate(row):
            fill_cell(table.cell(i, j), value, size=7.7)
    format_table(table, header=True, first_col=True, widths=[0.42, 2.12, 4.42], font_size=7.7)
    add_callout(
        doc,
        "Nguyên tắc gom sản phẩm",
        "Không biến mỗi module thành một nút rời. Các module dùng chung DeviceIdentity, Evidence, RiskGate, checkpoint, "
        "verify, rollback và timeline; giao diện bắt đầu từ mục tiêu người dùng, không từ tên công cụ kỹ thuật.",
        kind="decision",
    )


def add_priorities(doc: Document) -> None:
    add_section_title(doc, "9", "Ma trận ưu tiên")
    add_body(
        doc,
        "Thang 1–5; điểm là ước lượng chuyên gia theo 30% tác động người dùng, 25% giảm rủi ro, 20% độ phủ, "
        "15% khả thi với lõi hiện tại và 10% khác biệt. Cần hiệu chỉnh bằng phỏng vấn và telemetry opt-in.",
    )
    table = doc.add_table(rows=1 + len(FEATURES), cols=5)
    for j, h in enumerate(["#", "Hệ thống", "Điểm", "Pha", "Lý do xếp hạng"]):
        fill_cell(table.cell(0, j), h, bold=True, color=WHITE, size=8.2)
    rationales = {
        1: "Phòng sự cố trước khi người dùng bị mất quyền truy cập; tận dụng core read-only.",
        2: "Tần suất/thiệt hại cao; nối nhiều công cụ rời rạc thành một hành trình.",
        3: "Giảm random fixes và thời gian hỗ trợ; dùng timeline hiện có.",
        4: "Giá trị cao nhất lúc khẩn cấp nhưng cần test WinRE/hardware rộng.",
        5: "Bao phủ hầu hết PC và là nguyên nhân phổ biến sau update/reinstall.",
        6: "Khác biệt mạnh, rủi ro cao; chỉ mở dần theo adapter/model.",
        7: "Giảm vòng lặp hỗ trợ và rủi ro riêng tư.",
        8: "Giá trị dài hạn; độ phủ telemetry không đồng nhất.",
        9: "Rất hữu ích cho máy tự lắp; tốn dữ liệu model/manual.",
        10: "Bảo mật phải đi cùng tương thích; cần nhiều dependency checks.",
        11: "Tốt cho retention nhưng nhiều app/license ngoại lệ.",
        12: "Khác biệt, song cần dữ liệu OEM và kích thước/nguồn đáng tin.",
        13: "Hữu ích nhưng phụ thuộc phần cứng/dock/cáp và dữ liệu khó chuẩn hóa.",
        14: "Quan trọng về riêng tư nhưng tần suất thấp hơn các luồng P0/P1.",
    }
    for i, feature in enumerate(FEATURES, start=1):
        values = [str(feature["rank"]), feature["name"], feature["score"], feature["phase"], rationales[feature["rank"]]]
        for j, value in enumerate(values):
            fill_cell(table.cell(i, j), value, size=7.6)
    format_table(table, header=True, first_col=True, widths=[0.38, 2.25, 0.56, 0.48, 3.29], font_size=7.6)


def add_roadmap(doc: Document) -> None:
    add_section_title(doc, "10", "Lộ trình và launch gates")
    milestones = [
        ("M0", "0–8 tuần", "Evidence contracts, signed knowledge packs, privacy/redaction, durable resume tests", "Mọi P0 biểu diễn Unknown/source/freshness/confidence; plan hash bền vững"),
        ("M1", "2–4 tháng", "Read-only: Recovery scan, install preflight, driver vault, timeline, storage map và startup inventory", "95% scan không side effect; 100% kết luận có source/confidence"),
        ("M2", "4–7 tháng", "Smart Cleanup có quarantine, driver transactions, recovery ladder, reboot resume, freeze recorder", "Verify độc lập; BitLocker/power gates; undo/rollback drill đạt ngưỡng"),
        ("M3", "7–12 tháng", "WPR profiler, app repair, update guard, BIOS diff, WinRE handoff và no-POST pilot", "Model mismatch không vào write; playbook có test thật hoặc guided-only"),
        ("M4", ">12 tháng", "Sandbox app, storage migration, peripherals, lifecycle và advisory", "Mở rộng theo usability/telemetry opt-in tại Việt Nam"),
    ]
    table = doc.add_table(rows=1 + len(milestones), cols=4)
    for j, h in enumerate(["Mốc", "Thời gian", "Deliverable", "Exit criteria"]):
        fill_cell(table.cell(0, j), h, bold=True, color=WHITE, size=8.5)
    for i, row in enumerate(milestones, start=1):
        for j, value in enumerate(row):
            fill_cell(table.cell(i, j), value, size=8.0)
    format_table(table, header=True, first_col=True, widths=[0.55, 0.82, 3.2, 2.39], font_size=8.0)

    metrics_heading = add_body(doc, "Chỉ số và gate", bold=True, color=SLATE, after=3)
    metrics_heading.paragraph_format.keep_with_next = True
    metrics = [
        ("Recovery path trước side effect rủi ro", "100%", "Bắt buộc"),
        ("Kết luận có source + freshness + confidence", "100%", "Bắt buộc"),
        ("False-success sau verify", "<0,5% pilot/test", "Bắt buộc"),
        ("Resume sau reboot/crash", "≥99,5% luồng hỗ trợ", "Bắt buộc"),
        ("Secret/PII trong bundle mặc định", "0 trong corpus", "Bắt buộc"),
        ("Model mismatch dẫn tới firmware write", "0", "Bắt buộc"),
        ("Hoàn tất P0 không cần terminal", "≥85% usability test", "Khuyến nghị"),
        ("Thời gian tìm nguyên nhân sơ bộ", "Giảm ≥50%", "Khuyến nghị"),
    ]
    table2 = doc.add_table(rows=1 + len(metrics), cols=3)
    for j, h in enumerate(["Signal", "Mục tiêu ban đầu", "Launch gate"]):
        fill_cell(table2.cell(0, j), h, bold=True, color=WHITE, size=8.4)
    for i, row in enumerate(metrics, start=1):
        for j, value in enumerate(row):
            fill_cell(table2.cell(i, j), value, size=8.1)
    format_table(table2, header=True, first_col=True, widths=[3.32, 1.78, 1.86], font_size=8.1)


def add_security_alternatives(doc: Document) -> None:
    add_section_title(doc, "11", "Bảo mật, riêng tư và giới hạn")
    for text in [
        "Tách quyền đọc, lập kế hoạch, thực thi và firmware write; mặc định least privilege và local-first.",
        "Secret-bearing evidence có classification riêng, không vào log/timeline; support bundle có preview, consent và retention.",
        "Knowledge pack/adapters có chữ ký, version, source URL, expiry và allowlist model/revision; hết hạn hạ support tier.",
        "Dump bộ nhớ có thể chứa nội dung người dùng nên luôn opt-in riêng; remote control chỉ sau cảnh báo trust/scam.",
        "Không áp security baseline doanh nghiệp mù quáng cho consumer; luôn so sánh lợi ích bảo vệ và rủi ro tương thích.",
    ]:
        add_bullet(doc, text)
    add_citations(doc, ["S21", "S32", "S33", "S34", "S35"])

    add_body(doc, "Giới hạn nghiên cứu", bold=True, color=SLATE, after=3)
    limits = [
        "Chưa phỏng vấn người dùng Việt Nam, cửa hàng lắp máy hoặc kỹ thuật viên; điểm ưu tiên là suy luận có cấu trúc.",
        "Tài liệu OEM thay đổi theo model, khu vực và BIOS revision; cần quy trình revalidation liên tục.",
        "Quick Machine Recovery, Point-in-time restore và Secure Boot transition đang thay đổi trong 2026; phải feature-detect.",
        "Telemetry NVMe/WHEA/battery/fan không có trên mọi máy; Unknown không được quy thành “ổn”.",
        "Auto mode cần thử phần cứng thật, threat model và usability test; báo cáo này chưa thay thế các bước đó.",
    ]
    for text in limits:
        add_bullet(doc, text)

    add_section_title(doc, "12", "Các lựa chọn đã loại")
    alternatives = [
        ("AI tự sửa mọi lỗi", "Dễ tạo cảm giác đơn giản", "Không có ground truth/authority; có thể bịa model, lệnh và nguyên nhân"),
        ("Universal BIOS writer", "Bao phủ rộng trên giấy", "Nguy cơ brick/BitLocker; API, rollback và recovery phụ thuộc OEM"),
        ("One-click optimize / update all", "Dễ marketing", "Nhiều biến cùng lúc, outcome mơ hồ, khó quy nguyên nhân/rollback"),
        ("Kho hàng trăm repair scripts", "Ra tính năng nhanh", "Thiếu state, dependency, verify và provenance"),
        ("Luôn cloud-first", "Dễ cập nhật tri thức", "Máy hỏng có thể mất mạng; log/dump nhạy cảm; cần local/offline path"),
    ]
    table = doc.add_table(rows=1 + len(alternatives), cols=3)
    for j, h in enumerate(["Lựa chọn", "Điểm hấp dẫn", "Lý do không chọn"]):
        fill_cell(table.cell(0, j), h, bold=True, color=WHITE, size=8.5)
    for i, row in enumerate(alternatives, start=1):
        for j, value in enumerate(row):
            fill_cell(table.cell(i, j), value, size=8.2)
    format_table(table, header=True, first_col=True, widths=[1.45, 1.88, 3.63], font_size=8.2)


def add_open_questions_decision(doc: Document) -> None:
    add_section_title(doc, "13", "Câu hỏi mở")
    questions = [
        "Nhóm 10–20 model/motherboard nào đại diện tốt nhất cho người dùng Việt Nam ở pilot?",
        "PC Orbit sẽ phân phối và thu hồi signed knowledge packs bằng kênh nào; ai chịu trách nhiệm xác minh OEM changes?",
        "Crash dump và telemetry nào được xử lý hoàn toàn local, và ngưỡng nào mới đề nghị upload?",
        "Cơ chế phone mode/offline handoff nào hoạt động khi PC không boot hoặc không có mạng?",
        "Tiêu chí phần cứng thật nào buộc Auto adapter phải đạt trước khi mở rộng rollout?",
    ]
    for q in questions:
        add_numbered(doc, q)

    add_section_title(doc, "14", "Quyết định và bước tiếp theo")
    add_callout(
        doc,
        "Phê duyệt hướng sản phẩm",
        "Xây PC Orbit như một lớp điều khiển tự phục vụ an toàn trên toàn vòng đời PC. Mốc đầu là vertical slice read-only: "
        "Recovery Readiness + Install Preflight + Driver Vault + Root-cause Timeline + Storage Map/Startup Inventory; "
        "sau đó mở Smart Cleanup có quarantine và Freeze Recorder.",
        kind="decision",
    )
    next_steps = [
        ("1", "Đóng schema M0", "DeviceIdentity, Evidence, RiskGate, RecoveryAsset, Hypothesis, SupportBundleManifest"),
        ("2", "Dựng corpus", "Máy OEM + desktop tự lắp; log install/driver/boot; PII redaction test set"),
        ("3", "Pilot read-only", "50–100 máy tự nguyện; đo completion, false-success và time-to-hypothesis"),
        ("4", "Mở giao dịch", "Driver rollback và recovery ladder trước; firmware chỉ 1 adapter/model set đã test"),
        ("5", "Mở rộng model", "Theo nhu cầu thực, không theo số lượng setting; công bố support tier rõ ràng"),
    ]
    table = doc.add_table(rows=1 + len(next_steps), cols=3)
    for j, h in enumerate(["Bước", "Deliverable", "Nội dung"]):
        fill_cell(table.cell(0, j), h, bold=True, color=WHITE, size=8.5)
    for i, row in enumerate(next_steps, start=1):
        for j, value in enumerate(row):
            fill_cell(table.cell(i, j), value, size=8.2)
    format_table(table, header=True, first_col=True, widths=[0.55, 1.62, 4.79], font_size=8.2)
    add_body(
        doc,
        "Tiêu chuẩn định vị: không quảng bá “hỗ trợ BIOS” theo số setting; quảng bá theo số model/workflow có đường "
        "xác minh và cứu hộ đã thử thật. Đây là ranh giới giữa công cụ đáng tin và một bộ mẹo có thể làm hỏng máy.",
        bold=True,
        color=NAVY,
        after=4,
    )


def add_evidence(doc: Document) -> None:
    add_section_title(doc, "15", "Claim-to-source ledger", page_break=True)
    add_body(
        doc,
        "Ledger dưới đây nối từng claim có ảnh hưởng đến thiết kế với nguồn gốc chính. Các link là tài liệu chính thức; "
        "điểm ưu tiên và lộ trình vẫn là suy luận của nghiên cứu này.",
    )
    table = doc.add_table(rows=1 + len(CLAIMS), cols=3)
    for j, h in enumerate(["ID", "Claim", "Nguồn"]):
        fill_cell(table.cell(0, j), h, bold=True, color=WHITE, size=8.2)
    for i, (claim_id, claim, source_ids) in enumerate(CLAIMS, start=1):
        fill_cell(table.cell(i, 0), claim_id, bold=True, color=NAVY, size=7.6)
        fill_cell(table.cell(i, 1), claim, size=7.6)
        cell = table.cell(i, 2)
        cell.text = ""
        p = cell.paragraphs[0]
        for idx, sid in enumerate(source_ids):
            title, url = SOURCES[sid]
            add_hyperlink(p, f"[{sid}] {title}", url, color=BLUE, underline=False)
            if idx != len(source_ids) - 1:
                p.add_run("\n")
    format_table(table, header=True, first_col=True, widths=[0.48, 3.77, 2.71], font_size=7.6)

    add_section_title(doc, "16", "Nguồn tham khảo")
    for sid, (title, url) in SOURCES.items():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{sid}  ")
        style_run(r, bold=True, color=NAVY, size=8.1)
        add_hyperlink(p, title, url, color=BLUE, underline=True)
        for run in p.runs:
            if run.text and not run.text.startswith(sid):
                style_run(run, size=8.1)


def add_document_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "PC Orbit — Trợ lý tự phục vụ PC an toàn"
    props.subject = "Deep research: chức năng chuyên sâu cho cài máy, BIOS/UEFI, chẩn đoán, cứu hộ và vòng đời PC"
    props.author = "OpenAI — Deep Research"
    props.keywords = "PC Orbit, Windows, BIOS, UEFI, recovery, drivers, installation, troubleshooting"
    props.comments = "Research snapshot as of 2026-09-04; official primary sources linked in the document."


def remove_unused_template_links(path: Path) -> None:
    """Remove external relationships left behind by cleared template placeholders."""
    rewritten = path.with_name(path.stem + ".rewritten.docx")
    rel_path = "word/_rels/document.xml.rels"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    with ZipFile(path, "r") as source_zip, ZipFile(rewritten, "w", ZIP_DEFLATED) as target_zip:
        for info in source_zip.infolist():
            payload = source_zip.read(info.filename)
            if info.filename == rel_path:
                root = ET.fromstring(payload)
                for rel in list(root):
                    target = rel.attrib.get("Target", "")
                    if rel.attrib.get("TargetMode") == "External" and "example.com" in target:
                        root.remove(rel)
                ET.register_namespace("", rel_ns)
                payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target_zip.writestr(info, payload)
    rewritten.replace(path)


def build() -> None:
    doc = Document(str(REFERENCE))
    clear_document_body(doc)
    tune_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.70)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.70)
    section.right_margin = Inches(0.70)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    configure_footer(doc)
    add_document_properties(doc)

    add_cover(doc)
    add_executive_summary(doc)
    add_scope_goals(doc)
    add_problem_and_lifecycle(doc)
    add_architecture(doc)
    add_safety_contracts(doc)
    add_feature_catalog(doc)
    add_expansion_catalog(doc)
    add_priorities(doc)
    add_roadmap(doc)
    add_security_alternatives(doc)
    add_open_questions_decision(doc)
    add_evidence(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    remove_unused_template_links(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
